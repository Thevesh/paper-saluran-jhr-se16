"""
Saluran-level results and demographics, as a formatted Excel workbook.
Usage: python3 saluran_xlsx.py        -> data/jhr_se16_saluran.xlsx

One sheet per election (ordinary polling-day salurans; the early and postal bloc is excluded, as
everywhere in this repo), plus a Notes sheet. Every percentage is a live formula over the raw
counts, which sit in a collapsed "Counts" section at the right of each sheet.

  Turnout      ballots issued / registered
  BN, PH, ...  % of valid votes (BERSAMA broken out at SE-16 only, as in the estimates)
  Malay, ...   % of the saluran's registered electorate, from the voter roll
  18-29, ...   likewise, by age band
"""

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import common

OUT = common.OUT
FONT = "Arial"
# section colours: (band fill, header fill) -- a strong band, a pale header row beneath it
COLOURS = {
    "Location": ("595959", "E7E6E6"),
    "Turnout & results": ("1F4E79", "DDEBF7"),
    "Ethnicity (% of registered)": ("375623", "E2EFDA"),
    "Age (% of registered)": ("833C0B", "FCE4D6"),
    "Counts": ("7F7F7F", "F2F2F2"),
}
PARTY_LABEL = {"BERSAMA": "BERSAMA", "OTHER": "Other"}
ETH_LABEL = {"Other": "Other eth."}


def parties(el):
    return [c for c in common.choices(el) if c not in ("REJECTED", "NOTRETURNED", "NONVOTE")]


def sheet(wb, el):
    df = common.load(el).copy()
    df["ballots"] = df.V - df.NONVOTE
    P = parties(el)
    ws = wb.create_sheet(common.cfg(el)["short"])

    # ---- column plan: (section, header, kind, source)
    loc = [("Location", h, "text", c) for h, c in
           [("State", None), ("Parlimen", "parlimen"), ("DUN", "dun"), ("DM", "dm"),
            ("PM", "pm"), ("Saluran", "saluran")]]
    counts = ([("Counts", "Registered", "count", "V"), ("Counts", "Ballots issued", "count", "ballots"),
               ("Counts", "Valid votes", "count", "valid")]
              + [("Counts", PARTY_LABEL.get(p, p), "count", p) for p in P]
              + [("Counts", f"{ETH_LABEL.get(g, g)} (n)", "count", f"n_{g}") for g in common.ETH]
              + [("Counts", f"{g} (n)", "count", f"n_{g}") for g in common.AGE])
    res = ([("Turnout & results", "Turnout", "pct", ("ballots", "V"))]
           + [("Turnout & results", f"% {PARTY_LABEL.get(p, p)}", "pct", (p, "valid")) for p in P])
    eth = [("Ethnicity (% of registered)", f"% {ETH_LABEL.get(g, g)}", "pct", (f"n_{g}", "V"))
           for g in common.ETH]
    age = [("Age (% of registered)", f"% {g}", "pct", (f"n_{g}", "V")) for g in common.AGE]
    cols = loc + res + eth + age + counts
    letter = {}                      # count source -> column letter, for the formulas
    for j, (sec, h, kind, src) in enumerate(cols, start=1):
        if kind == "count":
            letter[src] = get_column_letter(j)
    # valid votes is itself a formula over the party counts
    party_letters = [letter[p] for p in P]

    # ---- header: band row (merged per section) + column names
    thin = Side(style="thin", color="BFBFBF")
    j = 1
    while j <= len(cols):
        sec = cols[j - 1][0]
        k = j
        while k + 1 <= len(cols) and cols[k][0] == sec:
            k += 1
        ws.merge_cells(start_row=1, start_column=j, end_row=1, end_column=k)
        c = ws.cell(row=1, column=j, value=sec)
        c.fill = PatternFill("solid", fgColor=COLOURS[sec][0])
        c.font = Font(name=FONT, bold=True, color="FFFFFF", size=11)
        c.alignment = Alignment(horizontal="center", vertical="center")
        j = k + 1
    for j, (sec, h, kind, src) in enumerate(cols, start=1):
        c = ws.cell(row=2, column=j, value=h)
        c.fill = PatternFill("solid", fgColor=COLOURS[sec][1])
        c.font = Font(name=FONT, bold=True, size=10)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = Border(bottom=Side(style="medium", color=COLOURS[sec][0]))
    ws.row_dimensions[1].height = 20
    ws.row_dimensions[2].height = 30

    # ---- body
    df = df.sort_values(["parlimen", "dun", "dm", "saluran"]).reset_index(drop=True)
    body = Font(name=FONT, size=10)
    for i, rd in enumerate(df.to_dict("records"), start=3):
        for j, (sec, h, kind, src) in enumerate(cols, start=1):
            if kind == "text":
                v = common.STATE if src is None else rd[src]
                c = ws.cell(row=i, column=j, value=int(v) if src == "saluran" else v)
            elif kind == "count":
                if src == "valid":
                    c = ws.cell(row=i, column=j,
                                value="=" + "+".join(f"{L}{i}" for L in party_letters))
                else:
                    c = ws.cell(row=i, column=j, value=int(rd[src]))
                c.number_format = "#,##0"
            else:
                num, den = src
                n, d = f"{letter[num]}{i}", f"{letter[den]}{i}"
                c = ws.cell(row=i, column=j, value=f'=IF({d}>0,{n}/{d},"")')
                c.number_format = "0.0%"
            c.font = body
            if sec != "Location":
                c.alignment = Alignment(horizontal="right")
    last = len(df) + 2

    # ---- sheet furniture
    widths = {"State": 15, "Parlimen": 20, "DUN": 20, "DM": 28, "PM": 34, "Saluran": 8}
    for j, (sec, h, kind, src) in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(j)].width = widths.get(h, 11 if kind == "pct" else 12)
    first_count = get_column_letter(len(cols) - len(counts) + 1)
    last_col = get_column_letter(len(cols))
    ws.column_dimensions.group(first_count, last_col, hidden=True, outline_level=1)
    ws.freeze_panes = "G3"
    ws.auto_filter.ref = f"A2:{last_col}{last}"
    ws.sheet_view.zoomScale = 90
    # light banding on location cells for readability
    for i in range(3, last + 1, 2):
        for j in range(1, 7):
            ws.cell(row=i, column=j).fill = PatternFill("solid", fgColor="FAFAFA")
    return len(df)


def notes(wb, counts):
    ws = wb.create_sheet("Notes", 0)
    rows = [
        (f"{common.STATE}: saluran-level results and demographics", True),
        ("", False),
        ("One sheet per election: " + ", ".join(
            f"{common.cfg(el)['label']} ({counts[el]:,} salurans)" for el in common.ORDER) + ".", False),
        ("Ordinary polling-day salurans only. The early (/00) and postal (/UP) bloc is excluded, "
         "as in every estimate in this repository: postal votes have no roll electorate.", False),
        ("Turnout = ballots issued / registered electorate.", False),
        ("Party columns = % of valid votes. BERSAMA is shown separately at SE-16, where it "
         "contested alone; it did not exist at SE-15 or GE-15. 'Other' is every other valid "
         "vote.", False),
        ("Ethnicity and age columns = % of the saluran's registered electorate, counted exactly off "
         "the public voter roll. 'Other eth.' covers all other non-Malay, non-Chinese, "
         "non-Indian voters (Orang Asli, Bumiputera Sabah/Sarawak and others). Age = election year minus birth year.", False),
        ("Every percentage is a formula over the raw counts in the collapsed 'Counts' section at "
         "the right of each sheet (click the + above the column letters to expand it).", False),
        ("Sources: SPR saluran scoresheets (compiled in paper-meco-saluran) and the public voter "
         "roll (ElectionData.MY).", False),
    ]
    for i, (t, bold) in enumerate(rows, start=1):
        c = ws.cell(row=i, column=1, value=t)
        c.font = Font(name=FONT, bold=bold, size=13 if bold else 10)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 120
    r0 = len(rows) + 2
    ws.cell(row=r0, column=1, value="Column sections").font = Font(name=FONT, bold=True, size=10)
    for k, (sec, (band, light)) in enumerate(COLOURS.items(), start=1):
        c = ws.cell(row=r0 + k, column=1, value=sec)
        c.fill = PatternFill("solid", fgColor=band)
        c.font = Font(name=FONT, bold=True, color="FFFFFF", size=10)


def main():
    wb = Workbook()
    wb.remove(wb.active)
    counts = {el: sheet(wb, el) for el in common.ORDER}
    notes(wb, counts)
    wb.calculation.fullCalcOnLoad = True    # formulas are saved without cached values
    path = f"{common.REPO}/data/jhr_se16_saluran.xlsx"
    wb.save(path)
    print(f"wrote {path}: " + ", ".join(f"{common.cfg(e)['short']} {n:,} rows" for e, n in counts.items()))


if __name__ == "__main__":
    main()
